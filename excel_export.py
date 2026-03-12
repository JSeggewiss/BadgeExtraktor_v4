"""Excel-Export mit openpyxl – detaillierte Kontaktliste"""

from datetime import datetime
from pathlib import Path
from typing import List, Dict

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


COLUMNS = [
    "Nachname",
    "Vorname",
    "Institution",
    "Position/Funktion",
    "Ort",
    "E-Mail",
    "Quelle",  # Website / Profil- oder Team-URL
]

HEADER_FILL = PatternFill("solid", start_color="2E4057")   # Dunkelblau
HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri", size=11)
DATA_FONT   = Font(name="Calibri", size=10)


def _style_header(ws, header_row: int):
    for col_idx, col_name in enumerate(COLUMNS, start=2):
        cell = ws.cell(row=header_row, column=col_idx)
        cell.value = col_name
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[header_row].height = 22


def _set_col_widths(ws, data_rows):
    widths = {col: len(col) + 2 for col in COLUMNS}
    for row in data_rows:
        for col in COLUMNS:
            val = str(row.get(col, "") or "")
            if len(val) > widths[col]:
                widths[col] = len(val)
    for col_idx, col_name in enumerate(COLUMNS, start=2):
        w = min(max(widths[col_name] + 2, 12), 55)
        ws.column_dimensions[get_column_letter(col_idx)].width = w


def save_to_excel(records: List[Dict], output_path: str | Path) -> Path:
    """Speichert eine Liste von Badge-Dicts als formatierte Excel-Datei.
    Gibt den endgültigen Pfad zurück."""
    output_path = Path(output_path)

    wb = Workbook()
    ws = wb.active
    ws.title = "Badges"

    # Linke Marginalspalte
    ws.column_dimensions["A"].width = 3

    # Titel
    ws.row_dimensions[1].height = 10  # kleiner Abstand oben
    title_cell = ws.cell(row=2, column=2, value="Badge-Kontakte (detailliert)")
    title_cell.font = Font(name="Calibri", size=14, bold=True, color="2E4057")
    ws.merge_cells(f"B2:{get_column_letter(len(COLUMNS) + 1)}2")
    ws.row_dimensions[2].height = 24

    subtitle = ws.cell(
        row=3, column=2,
        value=f"Erstellt: {datetime.now().strftime('%d.%m.%Y %H:%M')}  |  {len(records)} Einträge"
    )
    subtitle.font = Font(name="Calibri", size=9, color="777777")
    ws.merge_cells(f"B3:{get_column_letter(len(COLUMNS) + 1)}3")
    ws.row_dimensions[3].height = 16

    HEADER_ROW = 5
    _style_header(ws, HEADER_ROW)

    # Daten eintragen
    for row_idx, record in enumerate(records, start=HEADER_ROW + 1):
        ws.row_dimensions[row_idx].height = 18
        for col_idx, col_name in enumerate(COLUMNS, start=2):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.value = record.get(col_name, "")
            cell.font = DATA_FONT
            cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)

            # Quell-URL als Hyperlink
            if col_name == "Quelle" and cell.value and str(cell.value).startswith("http"):
                cell.hyperlink = cell.value
                cell.value = "Link"
                cell.font = Font(name="Calibri", size=10, color="0563C1", underline="single")

    # Excel-Tabelle anlegen
    last_row = HEADER_ROW + len(records)
    last_col = get_column_letter(len(COLUMNS) + 1)
    tab = Table(
        displayName="BadgeTabelle",
        ref=f"B{HEADER_ROW}:{last_col}{last_row}",
    )
    tab.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    ws.add_table(tab)

    # Einfrieren unterhalb Header
    ws.freeze_panes = f"B{HEADER_ROW + 1}"

    _set_col_widths(ws, records)

    wb.save(output_path)
    return output_path


def load_records(path: str | Path) -> list[dict]:
    """Liest bestehende Badge-Datensätze aus einer Excel-Datei."""
    path = Path(path)
    if not path.exists():
        return []
    try:
        wb = load_workbook(path, data_only=True)
        ws = wb.active
        HEADER_ROW = 5
        records = []
        for row in ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True):
            if any(v for v in row[1:]):  # Spalte A leer lassen
                record = {}
                for col_idx, col_name in enumerate(COLUMNS):
                    record[col_name] = str(row[col_idx + 1] or "")
                records.append(record)
        return records
    except Exception:
        return []
